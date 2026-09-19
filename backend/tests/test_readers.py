"""
Tests for backend/ai/readers.py. All test files are generated on the fly: no API calls, no real data needed.

    python -m pytest backend/tests -v
"""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # backend/

from ai.readers import ReaderError, read_document  # noqa: E402


def make_pdf(lines) -> bytes:
    """A tiny hand-built PDF. With no lines it is a blank page, like a scanned image would look to a text reader."""
    content = "BT /F1 12 Tf 50 750 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET" if lines else ""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF".encode()
    return out


def test_txt_is_decoded():
    assert read_document("SHIPPING INSTRUCTION\nShipper: ACME".encode(), "x.txt") == "SHIPPING INSTRUCTION\nShipper: ACME"


def test_xlsx_rows_become_lines_and_blank_rows_and_cells_are_skipped():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "SI"
    sheet.append(["SHIPPING INSTRUCTION"])
    sheet.append([])
    sheet.append(["Shipper", "ACME SDN BHD"])
    sheet.append(["Containers", 3.0, None, "x 40'HC"])
    other = workbook.create_sheet("Notes")
    other.append(["Gross Weight (KG)", 22000])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = read_document(buffer.getvalue(), "a.xlsx")
    assert text.splitlines() == [
        "=== SHEET: SI ===",
        "SHIPPING INSTRUCTION",
        "Shipper | ACME SDN BHD",
        "Containers | 3 | x 40'HC",  # 3.0 is shown as 3
        "=== SHEET: Notes ===",
        "Gross Weight (KG) | 22000",
    ]


def test_docx_keeps_paragraphs_and_tables_in_order_and_merged_cells_once():
    import docx

    document = docx.Document()
    document.add_paragraph("BILL OF LADING (DRAFT)")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Shipper"
    merged = table.cell(0, 1).merge(table.cell(0, 2))
    merged.text = "ACME SDN BHD"
    table.cell(1, 0).text = "POL"
    table.cell(1, 1).text = "PORT KLANG"
    table.cell(1, 2).text = "(MYPKG)"
    document.add_paragraph("Freight: PREPAID")
    buffer = io.BytesIO()
    document.save(buffer)

    assert read_document(buffer.getvalue(), "b.docx").splitlines() == [
        "BILL OF LADING (DRAFT)",
        "Shipper | ACME SDN BHD",  # merged cell appears once, not twice
        "POL | PORT KLANG | (MYPKG)",
        "Freight: PREPAID",
    ]


def test_pdf_text_is_extracted():
    text = read_document(make_pdf(["SHIPPING INSTRUCTION", "Shipper: ACME SDN BHD"]), "c.pdf")
    assert "SHIPPING INSTRUCTION" in text
    assert "Shipper: ACME SDN BHD" in text


def test_pdf_without_a_text_layer_returns_empty_string():
    assert read_document(make_pdf([]), "scan.pdf") == ""


@pytest.mark.parametrize("name", ["broken.pdf", "broken.docx", "broken.xlsx"])
def test_corrupt_files_raise_reader_error_not_a_random_crash(name):
    with pytest.raises(ReaderError):
        read_document(b"this is not a real document", name)


def test_unsupported_type_raises_reader_error():
    with pytest.raises(ReaderError, match="unsupported"):
        read_document(b"x", "picture.png")