"""
readers.py - turn an attachment (.txt / .xlsx / .docx / .pdf) into plain text for extract_fields().

There is NO AI and NO API call in this file, so it is free to run and test.

    from ai.readers import read_document
    text = read_document(inbox.read_bytes("attachments/email_005_SI.xlsx"), "attachments/email_005_SI.xlsx")

- Tables are flattened to one line per row, cells separated by " | ".
- A scanned (image-only) PDF has no text layer: read_document() returns "" and the caller should treat the
  document as "needs OCR / vision" or "unreadable" (Role 5's Document AI, or Claude reading the PDF directly).
- A corrupt or password-protected file raises ReaderError (never a crash with a strange traceback).

Extra libraries:   pip install openpyxl python-docx pdfplumber
"""
import io
from pathlib import Path

SUPPORTED_SUFFIXES = (".txt", ".xlsx", ".xlsm", ".docx", ".pdf")


class ReaderError(Exception):
    """The file could not be turned into text (corrupt, protected, unsupported type, library missing)."""


def _need(module_name: str, pip_name: str):
    try:
        return __import__(module_name)
    except ImportError as err:
        raise ReaderError(f"the '{pip_name}' library is not installed - run: pip install openpyxl python-docx pdfplumber") from err


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _read_xlsx(data: bytes) -> str:
    openpyxl = _need("openpyxl", "openpyxl")
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)  # data_only: values, not formulas
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"=== SHEET: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            cells = [text for text in (_cell_text(v) for v in row) if text]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _read_docx(data: bytes) -> str:
    docx = _need("docx", "python-docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    lines = []
    # Walk the body in order so paragraphs and tables stay where they are in the document.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells, previous = [], None
                for cell in row.cells:
                    text = " ".join(cell.text.split())
                    if text and cell._tc is not previous:  # a merged cell shows up several times: keep it once
                        cells.append(text)
                    previous = cell._tc
                if cells:
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def _read_pdf(data: bytes) -> str:
    pdfplumber = _need("pdfplumber", "pdfplumber")
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"--- page {number} ---\n{text}")
    return "\n".join(pages)


def read_document(data: bytes, filename: str) -> str:
    """Plain text of an attachment. Raises ReaderError if it cannot be opened."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            return data.decode("utf-8", errors="replace")
        if suffix in (".xlsx", ".xlsm"):
            return _read_xlsx(data)
        if suffix == ".docx":
            return _read_docx(data)
        if suffix == ".pdf":
            return _read_pdf(data)
    except ReaderError:
        raise
    except Exception as err:  # corrupt, password-protected, wrong file type, ...
        raise ReaderError(f"could not open {Path(filename).name}: {type(err).__name__}: {err}") from err
    raise ReaderError(f"unsupported file type: {suffix or filename}")


def pdf_page_info(data: bytes):
    """(number of pages, number of embedded images) of a PDF, or None if it cannot be opened."""
    try:
        pdfplumber = _need("pdfplumber", "pdfplumber")
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return len(pdf.pages), sum(len(page.images) for page in pdf.pages)
    except Exception:
        return None
